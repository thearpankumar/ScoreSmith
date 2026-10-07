"""Pure end-to-end tests of the workbook renderer: build from a synthetic bundle, reopen with openpyxl."""

from __future__ import annotations

import uuid
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.reporting import palette as P
from app.reporting.export_data import EvidenceItem, ExportBundle, normalize_evidence
from app.reporting.insights import compute_insights
from app.reporting.workbook import ExportOptions, build_workbook
from tests.reporting_fixtures import make_bundle, make_eval, make_version

EVIL = ['=HYPERLINK("http://evil","x")', "+cmd|' /C calc'!A0", "@SUM(1+1)", "-2+3"]


def _open(bundle: ExportBundle, **opts):
    data = build_workbook(bundle, ExportOptions(generated_by="Tester", **opts))
    assert data[:2] == b"PK"
    return load_workbook(BytesIO(data))


def _texts(ws) -> list[str]:
    return [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


STUDENTS = ["01 Priya Shah", "02 Omar Ali", "03 Asha Menon", "04 Dev Patel", "05 Mia Chen", "06 Rahul Kumar"]


def test_sheet_order_and_summary_first() -> None:
    wb = _open(make_bundle())
    assert wb.sheetnames == ["Summary", "Leaderboard", "Evaluations", "KPI Matrix", "KPI Detail", *STUDENTS,
                             "KPI Reference", "Guidelines", "Notes"]
    assert wb.active.title == "Summary"


def test_summary_has_insights_above_every_table() -> None:
    ws = _open(make_bundle())["Summary"]
    assert ws["B1"].value == "Evaluation Report"
    rows = {str(c.value): c.row for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    take_header = rows["Key takeaways"]
    first_table_heading = min(r for t, r in rows.items() if t.startswith(("At a glance", "Score distribution", "Top performers")))
    sentences = [r for t, r in rows.items() if t.startswith("6 evaluations are included")]
    assert take_header < sentences[0] < first_table_heading
    assert take_header <= 6
    assert ws["B3"].value.startswith("Average 6.8 / 10 against a target of 7.0")  # verdict line right under the title
    txt = " ".join(_texts(ws))
    assert "Highest: Priya Shah at 9.1" in txt and "Needs attention" in txt and "Risk flags" in txt
    assert "How to read the colours" in txt and "Contents" in txt and "Student sheets (6)" in txt
    assert "Target 7.0  (default: no target set)" in txt and "Exceeds target" in txt


def test_leaderboard_order_ties_emails_and_links() -> None:
    ws = _open(make_bundle())["Leaderboard"]
    header_row = next(c.row for row in ws.iter_rows() for c in row if c.value == "Rank")
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 6)]
    assert [r[1] for r in rows] == ["Priya Shah", "Omar Ali", "Asha Menon", "Dev Patel", "Mia Chen", "Rahul Kumar"]
    assert [r[0] for r in rows] == [1, 2, 3, 3, 5, 6]
    assert ws.cell(header_row + 3, 1).number_format == '"T"0' and ws.cell(header_row + 2, 1).number_format == "0"
    assert ws.cell(header_row + 1, 3).value == "priya@example.com"
    assert ws.cell(header_row + 1, 3).hyperlink.target == "mailto:priya@example.com"
    assert ws.cell(header_row + 4, 3).value is None  # Dev Patel has no email -> blank, no link
    assert [r[3] for r in rows] == sorted((r[3] for r in rows), reverse=True)
    assert "tblLeaderboard" in ws.tables
    # the name links to that student's own sheet; the Band is target-relative (default target 7.0 here)
    assert ws.cell(header_row + 1, 2).hyperlink.location.strip("'").startswith("01 Priya Shah")
    assert [r[4] for r in rows] == ["Exceeds target", "Exceeds target", "Meets target", "Meets target",
                                    "Well below target", "Critical"]
    assert [r[5] for r in rows][0] == pytest.approx(2.1) and [r[5] for r in rows][-1] == pytest.approx(-3.6)


def test_single_evaluation_has_no_leaderboard_and_gets_a_student_sheet() -> None:
    wb = _open(make_bundle([("Solo", "solo@x.com", 8.3, [9, 8, 8, 8])]))
    assert "Leaderboard" not in wb.sheetnames
    assert "01 Solo" in wb.sheetnames and wb.sheetnames[0] == "Summary"
    assert any(t.startswith("Solo scored 8.3 / 10") for t in _texts(wb["Summary"]))
    assert wb["Summary"]["B3"].value.startswith("Solo: 8.3 / 10 — Exceeds target")


def _fill(cell) -> str:
    return cell.fill.fgColor.rgb[-6:].upper()


def test_score_cells_use_static_target_relative_fills_not_color_scales() -> None:
    wb = _open(make_bundle())
    for ws in wb.worksheets:
        assert not [r for rng in ws.conditional_formatting for r in rng.rules if r.type == "colorScale"], ws.title
    board = wb["Leaderboard"]
    head = next(c.row for row in board.iter_rows() for c in row if c.value == "Rank")
    for offset, (score, band_key) in enumerate([(9.1, "exceeds"), (8.2, "exceeds"), (7.5, "meets"), (7.5, "meets"),
                                                (5.2, "well_below"), (3.4, "critical")], start=1):
        band = P.TARGET_BAND_BY_KEY[band_key]
        assert board.cell(head + offset, 4).value == score
        assert _fill(board.cell(head + offset, 4)) == band.tint.lstrip("#")  # pale tint on the score
        assert _fill(board.cell(head + offset, 5)) == band.fill.lstrip("#")  # solid colour on the Band label


def test_a_low_target_changes_the_colours_not_just_the_numbers() -> None:
    """The same 4.2 is green for a target of 4 and dark red for a target of 9."""
    results = {}
    for target in (4.0, 9.0):
        wb = _open(make_bundle([("Pat", "p@x.com", 4.2, [4, 4, 4, 4]), ("Quinn", "q@x.com", 2.0, [2, 2, 2, 2])], target=target))
        board = wb["Leaderboard"]
        head = next(c.row for row in board.iter_rows() for c in row if c.value == "Rank")
        results[target] = (_fill(board.cell(head + 1, 4)), board.cell(head + 1, 5).value, board.cell(head + 1, 6).value)
    assert results[4.0][0] == P.TARGET_BAND_BY_KEY["meets"].tint.lstrip("#") and results[4.0][1] == "Meets target"
    assert results[9.0][0] == P.TARGET_BAND_BY_KEY["critical"].tint.lstrip("#") and results[9.0][1] == "Critical"
    assert results[4.0][2] == pytest.approx(0.2) and results[9.0][2] == pytest.approx(-4.8)
    wb = _open(make_bundle([("Pat", "p@x.com", 4.2, [4, 4, 4, 4]), ("Quinn", "q@x.com", 2.0, [2, 2, 2, 2])], target=4.0))
    txt = " ".join(_texts(wb["Summary"]))
    assert "Target 4.0" in txt and "default" not in txt.split("How to read the colours")[1].split("Contents")[0].split("Target 4.0")[0]
    assert "≥ 4.4" in txt  # the legend states the thresholds for THIS target


def test_no_formula_cells_and_injection_is_neutralised() -> None:
    vid, kpis = make_version()
    evals = [make_eval(vid, kpis, bad, f"{i}@x.com", 8.0 - i, [8, 8, 8, 8]) for i, bad in enumerate(EVIL)]
    evals.append(make_eval(vid, kpis, "Normal", "=1+1@x.com", 5.0, [5] * 4))
    wb = _open(ExportBundle(evals, {vid: kpis}), filter_summary="=cmd|' /C calc'!A0")
    names_seen = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                if c.value in EVIL:
                    names_seen.add(c.value)
                    assert c.data_type == "s"
                    assert c.quotePrefix
    assert names_seen == set(EVIL)
    board = wb["Leaderboard"]
    assert any(c.value == EVIL[0] and c.quotePrefix for row in board.iter_rows() for c in row)


def test_matrix_layout_and_average_row() -> None:
    ws = _open(make_bundle())["KPI Matrix"]
    assert ws.freeze_panes == "D6"
    assert "Clarity" in str(ws["D4"].value) and "Communication" in str(ws["D3"].value)
    assert ws["A5"].value == "Average (selected)"
    assert ws["A6"].value == "Priya Shah" and ws["C6"].value == 9.1
    assert _fill(ws["C6"]) == P.TARGET_BAND_BY_KEY["exceeds"].tint.lstrip("#")  # 9.1 vs target 7.0
    assert _fill(ws["D6"]) == P.TARGET_BAND_BY_KEY["exceeds"].tint.lstrip("#")


def test_include_reasoning_false_drops_text_columns() -> None:
    with_text = _open(make_bundle())["KPI Detail"]
    without = _open(make_bundle(), include_reasoning=False)["KPI Detail"]

    def headers(ws):
        row = next(r for r in ws.iter_rows() if any(c.value == "Evaluation" for c in r))
        return [c.value for c in row if c.value]

    assert {"Reasoning", "Evidence", "Matched guideline", "Evidence source"} <= set(headers(with_text))
    assert not ({"Reasoning", "Evidence", "Matched guideline", "Evidence source"} & set(headers(without)))
    # student sheets keep the score rows but lose the Why / Evidence rows
    sheet = _open(make_bundle(), include_reasoning=False)["01 Priya Shah"]
    labels = {str(c.value) for row in sheet.iter_rows(min_col=1, max_col=1) for c in row if c.value}
    assert "Clarity" in labels and not any(v.startswith(("Why", "Evidence")) for v in labels)


def test_mixed_scorecards_create_distinct_matrix_sheets_and_warning() -> None:
    v1, k1 = make_version()
    v2, k2 = make_version()
    evals = [
        make_eval(v1, k1, "A", "a@x.com", 9.0, [9] * 4, scorecard="Scorecard / One: [x]"),
        make_eval(v2, k2, "B", "b@x.com", 6.0, [6] * 4, scorecard="Scorecard / One: [x]!!"),
    ]
    wb = _open(ExportBundle(evals, {v1: k1, v2: k2}))
    matrices = [n for n in wb.sheetnames if n.startswith("Matrix")]
    assert len(matrices) == 2 and len({m.lower() for m in matrices}) == 2
    assert all(len(n) <= 31 and not any(ch in n for ch in "[]:*?/\\") for n in wb.sheetnames)
    assert any("different scorecards" in t for t in _texts(wb["Leaderboard"]))


def test_failed_and_unscored_evaluations_are_never_exported() -> None:
    vid, kpis = make_version()
    evals = [
        make_eval(vid, kpis, "Done", "d@x.com", 8.0, [8] * 4),
        make_eval(vid, kpis, "Done2", "d2@x.com", 7.0, [7] * 4),
        make_eval(vid, kpis, "Broken", "b@x.com", None, [], status="failed"),
        make_eval(vid, kpis, "Waiting", "w@x.com", 9.9, [9] * 4, status="queued"),
    ]
    bundle = ExportBundle(evals, {vid: kpis}, missing_ids=[uuid.uuid4()], excluded_count=1)
    wb = _open(bundle)
    for ws in wb.worksheets:
        assert not any("Broken" in t or "Waiting" in t for t in _texts(ws)), ws.title
    assert not any(n.endswith(("Broken", "Waiting")) for n in wb.sheetnames)
    assert wb.sheetnames.count("Evaluations") == 1 and len([n for n in wb.sheetnames if n[:2].isdigit()]) == 2
    notes = " ".join(_texts(wb["Notes"]))
    assert "4 selected evaluations were not exported (deleted, failed or not yet completed)" in notes  # 1 gone + 1 + 2
    ev_sheet = wb["Evaluations"]
    heads = [c.value for c in next(r for r in ev_sheet.iter_rows() if any(c.value == "Evaluation" for c in r))]
    assert "Status" not in heads and "Error" not in heads
    assert [r.ev.subject_name for r in compute_insights(bundle).scored] == ["Done", "Done2"]


def test_every_exported_student_gets_a_sheet_in_rank_order() -> None:
    wb = _open(make_bundle())
    assert [n for n in wb.sheetnames if n[:2].isdigit()] == STUDENTS
    board = wb["Leaderboard"]
    head = next(c.row for row in board.iter_rows() for c in row if c.value == "Rank")
    names = [board.cell(head + i, 2).value for i in range(1, 7)]
    assert names == [n[3:] for n in STUDENTS]


@pytest.mark.parametrize("n", [1, 40])
def test_builds_for_small_and_larger_selections(n: int) -> None:
    rows = [(f"Person {i}", f"p{i}@x.com", round(10 - (i * 10 / max(n, 1)), 2), [8, 7, 6, 5]) for i in range(n)]
    wb = _open(make_bundle(rows))
    assert wb.sheetnames[0] == "Summary"
    assert len([x for x in wb.sheetnames if x[:2].isdigit()]) == n  # one student sheet per exported evaluation


# ---- the HTTP layer, with the DB loader stubbed out (runs without Postgres) --------------------------------------
def test_endpoint_with_stubbed_loader(monkeypatch) -> None:
    import uuid

    from fastapi.testclient import TestClient

    import app.api.v1.evaluations_export as mod
    from app.config import get_settings
    from app.db import get_db
    from app.deps import get_current_user
    from app.main import app
    from app.models.user import User

    bundle = make_bundle()
    bundle.missing_ids = [uuid.uuid4()]

    async def fake_loader(db, ids):
        return bundle

    async def fake_db():
        yield None

    monkeypatch.setattr(mod, "load_export_bundle", fake_loader)
    app.dependency_overrides[get_db] = fake_db
    app.dependency_overrides[get_current_user] = lambda: User(id=uuid.uuid4(), email="t@x.com", name="Tester")
    try:
        ids = [str(e.id) for e in bundle.evaluations]
        saved = list(bundle.evaluations)
        resp = TestClient(app).post("/api/v1/evaluations/export", json={"evaluation_ids": ids, "filter_summary": "Status: Completed"})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert resp.headers["x-export-count"] == "6" and resp.headers["x-export-skipped"] == "1"
        assert 'filename="evaluations_Interview-Scorecard_' in resp.headers["content-disposition"]
        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames[0] == "Summary" and "Status: Completed" in " ".join(_texts(wb["Summary"]))
        # an export of evaluations that all lack a score is rejected
        for e in bundle.evaluations:
            e.status = "queued"
        assert TestClient(app).post("/api/v1/evaluations/export", json={"evaluation_ids": ids}).status_code == 422
        bundle.evaluations = []
        assert TestClient(app).post("/api/v1/evaluations/export", json={"evaluation_ids": ids}).status_code == 404
        bundle.excluded_count = 3  # found, but every one was failed / not completed -> 422, not 404
        assert TestClient(app).post("/api/v1/evaluations/export", json={"evaluation_ids": ids}).status_code == 422
        bundle.excluded_count = 0
        bundle.evaluations = list(saved)
        origin = get_settings().cors_origin_list[0]
        cors = TestClient(app).post("/api/v1/evaluations/export", json={"evaluation_ids": ids}, headers={"Origin": origin})
        assert "X-Export-Skipped" in cors.headers.get("access-control-expose-headers", "")
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)


# ------------------------------------------------------------------------------- KPI Reference + Guidelines
def _rich_version(*, level0: bool = False):
    """Section A (KPI one, nested sub-section 'Deep' holding KPI two), Section B (info-only KPI three)."""
    import uuid

    from app.reporting.export_data import ExportKpi

    vid = uuid.uuid4()
    a, deep, b = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()

    def leaf(name, parent, path, weight, included=True, guidelines=None):
        return ExportKpi(id=uuid.uuid4(), parent_id=parent, name=name, level=len(path), weight=weight,
                         included=included, is_leaf=True, order=0, path=path, category=path[0], guidelines=guidelines or {})

    full = {lv: f"Rubric for level {lv}" for lv in range(1, 11)}
    if level0:
        full[0] = "Nothing shown"
    kpis = [
        ExportKpi(id=a, parent_id=None, name="Section A", level=1, weight=None, included=True, is_leaf=False, order=0,
                  path=("Section A",)),
        leaf("KPI one", a, ("Section A", "KPI one"), 60.0, guidelines={**full, 3: '=HYPERLINK("http://evil","x")', 7: ""}),
        ExportKpi(id=deep, parent_id=a, name="Deep", level=2, weight=None, included=True, is_leaf=False, order=2,
                  path=("Section A", "Deep")),
        leaf("KPI two", deep, ("Section A", "Deep", "KPI two"), 40.0, guidelines={1: "Only level one", 10: "Only ten"}),
        ExportKpi(id=b, parent_id=None, name="Section B", level=1, weight=None, included=True, is_leaf=False, order=4,
                  path=("Section B",)),
        leaf("KPI three", b, ("Section B", "KPI three"), None, included=False),
    ]
    for i, k in enumerate(kpis):
        k.order = i
    return vid, kpis


def _rich_bundle(*, level0: bool = False):
    vid, kpis = _rich_version(level0=level0)
    n_leaves = sum(1 for k in kpis if k.is_leaf)
    evs = [make_eval(vid, kpis, "Ann", "ann@example.com", 8.0, [8, 6, 5]),
           make_eval(vid, kpis, "Bob", "bob@example.com", 6.0, [6, 4, 3])]
    for e in evs:
        e.kpis_total = n_leaves
    return ExportBundle(evaluations=evs, kpis_by_version={vid: kpis})


def _grid(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_reference_sheets_come_after_detail_before_notes_for_any_selection() -> None:
    for bundle in (_rich_bundle(), make_bundle([("Solo", "solo@example.com", 7.0, [7, 7, 7, 7])])):
        names = _open(bundle).sheetnames
        assert names[-3:] == ["KPI Reference", "Guidelines", "Notes"]
        assert names.index("KPI Reference") > names.index("Evaluations")


def test_reference_sheets_present_even_without_reasoning() -> None:
    wb = _open(_rich_bundle(), include_reasoning=False)
    assert "KPI Reference" in wb.sheetnames and "Guidelines" in wb.sheetnames
    grid = _grid(wb["Guidelines"])
    assert any(r[3] == "KPI one" and r[4 + 4] == "Rubric for level 5" for r in grid)  # the flag does not hide rubrics


def test_kpi_reference_hierarchy_weights_and_stats() -> None:
    ws = _open(_rich_bundle())["KPI Reference"]
    grid = _grid(ws)
    header = next(r for r in grid if r[0] == "Scorecard")
    assert header[:8] == ["Scorecard", "Version", "Type", "Section", "Sub-section", "KPI", "Weight (%)",
                          "Counts toward overall"]
    body = [r for r in grid if r[2] in ("Section", "Sub-section", "KPI")]
    assert [(r[2], r[3], r[4], r[5]) for r in body] == [
        ("Section", "Section A", None, None),
        ("KPI", "Section A", None, "KPI one"),
        ("Sub-section", "Section A", "Deep", None),
        ("KPI", "Section A", "Deep", "KPI two"),
        ("Section", "Section B", None, None),
        ("KPI", "Section B", None, "KPI three"),
    ]
    weights = {(r[2], r[3], r[5]): r[6] for r in body}
    assert weights[("Section", "Section A", None)] == 100.0  # 60 + 40 under A (including the nested sub-section)
    assert weights[("Sub-section", "Section A", None)] == 40.0
    assert weights[("KPI", "Section A", "KPI one")] == 60.0
    assert weights[("KPI", "Section B", "KPI three")] is None
    counts = {r[5]: r[7] for r in body if r[2] == "KPI"}
    assert counts == {"KPI one": "Yes", "KPI two": "Yes", "KPI three": "No (info only)"}
    stats = {r[5]: (r[9], r[10]) for r in body if r[2] == "KPI"}
    assert stats["KPI one"] == (2, 7.0) and stats["KPI two"] == (2, 5.0)  # mean of (8, 6) and (6, 4)
    assert any(isinstance(r[0], str) and r[0].startswith("Scorecard: Interview Scorecard") and "Version 1" in r[0]
               for r in grid)
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref


def test_kpi_reference_indents_nested_kpis() -> None:
    ws = _open(_rich_bundle())["KPI Reference"]
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value in ("KPI one", "KPI two", "KPI three")}
    assert cells["KPI one"].alignment.indent == 1 and cells["KPI two"].alignment.indent == 2


def test_guidelines_levels_1_to_10_columns_blanks_and_colours() -> None:
    ws = _open(_rich_bundle())["Guidelines"]
    grid = _grid(ws)
    header = next(r for r in grid if r[0] == "Scorecard")
    assert [str(h).split("\n")[0] for h in header[4:]] == [f"Level {i}" for i in range(1, 11)]  # no level 0 in the data
    assert "no KPI defines a level-0 rubric" in " ".join(str(c) for r in grid[:3] for c in r if c)
    one = next(r for r in grid if r[3] == "KPI one")
    assert one[4] == "Rubric for level 1" and one[4 + 9] == "Rubric for level 10" and one[4 + 4] == "Rubric for level 5"
    assert one[4 + 6] is None  # level 7 defined but empty -> blank
    two = next(r for r in grid if r[3] == "KPI two")
    assert two[4] == "Only level one" and two[13] == "Only ten" and two[4 + 4] is None
    three = next(r for r in grid if r[3] == "KPI three")
    assert all(v is None for v in three[4:])  # a KPI with no rubric at all still gets a row
    hdr_row = next(c.row for row in ws.iter_rows() for c in row if c.value == "Scorecard")
    low, high = ws.cell(hdr_row, 5).fill.fgColor.rgb[-6:], ws.cell(hdr_row, 14).fill.fgColor.rgb[-6:]
    assert int(low[:2], 16) > int(low[2:4], 16) and int(high[2:4], 16) > int(high[:2], 16)  # red-ish -> green-ish
    row_two = next(c.row for row in ws.iter_rows() for c in row if c.value == "KPI two")
    assert ws.cell(row_two, 4 + 5).fill.fgColor.rgb[-6:] == "EEEEEE"
    assert ws.freeze_panes == "E5" and ws.auto_filter.ref


def test_guidelines_adds_level_zero_column_only_when_data_has_it() -> None:
    ws = _open(_rich_bundle(level0=True))["Guidelines"]
    grid = _grid(ws)
    header = next(r for r in grid if r[0] == "Scorecard")
    assert str(header[4]).startswith("Level 0") and str(header[5]).startswith("Level 1") and len(header) == 15
    one = next(r for r in grid if r[3] == "KPI one")
    assert one[4] == "Nothing shown" and one[5] == "Rubric for level 1"
    assert "Level 0 column is included" in " ".join(str(c) for r in grid[:3] for c in r if c)


def test_reference_sheets_are_injection_safe_and_formula_free() -> None:
    wb = _open(_rich_bundle())
    cell = next(c for row in wb["Guidelines"].iter_rows() for c in row if c.value == '=HYPERLINK("http://evil","x")')
    assert cell.data_type == "s" and cell.quotePrefix
    for ws in wb.worksheets:
        assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row), ws.title


def test_guidelines_overlong_text_is_truncated_within_excel_row_limit() -> None:
    bundle = _rich_bundle()
    kpi = next(k for ks in bundle.kpis_by_version.values() for k in ks if k.name == "KPI one")
    kpi.guidelines[2] = "word " * 1600  # 8000 chars: far more than fits in a 409pt row at this width
    ws = _open(bundle)["Guidelines"]
    cell = next(c for row in ws.iter_rows() for c in row
                if c.column == 6 and isinstance(c.value, str) and c.value.startswith("word"))
    assert cell.value.endswith("…") and len(cell.value) < 8000
    assert ws.row_dimensions[cell.row].height <= 409


def test_reference_sheets_group_each_scorecard_version_separately() -> None:
    vid1, k1 = _rich_version()
    vid2, k2 = make_version()
    e1 = make_eval(vid1, k1, "Ann", "ann@example.com", 8.0, [8, 6, 5])
    e2 = make_eval(vid2, k2, "Bob", "bob@example.com", 6.0, [6, 6, 6, 6], scorecard="Other Card")
    e2.version_number = 3
    wb = _open(ExportBundle(evaluations=[e1, e2], kpis_by_version={vid1: k1, vid2: k2}))
    ref = _grid(wb["KPI Reference"])
    banners = [r[0] for r in ref if isinstance(r[0], str) and r[0].startswith("Scorecard: ")]
    assert len(banners) == 2  # sorted by scorecard name: "Interview Scorecard" before "Other Card"
    assert "Interview Scorecard" in banners[0] and "Version 1" in banners[0]
    assert "Other Card" in banners[1] and "Version 3" in banners[1]
    gl = _grid(wb["Guidelines"])
    assert {r[1] for r in gl if r[3] in ("Clarity", "KPI one")} == {"v1", "v3"}


# ------------------------------------------------------------------------------------------ student sheets
def _evidence_bundle(**kw):
    vid, kpis = make_version()
    ev_items = [EvidenceItem("We shipped the pilot in six weeks.", "Section 2"),
                EvidenceItem("Latency dropped from 800ms to 120ms.", "Appendix B"),
                EvidenceItem("=HYPERLINK(\"http://evil\",\"x\")", "")]
    a = make_eval(vid, kpis, "Ann Student", "ann@example.com", 8.0, [8, 6, 5, 7], evidence=ev_items, **kw)
    b = make_eval(vid, kpis, "Bob Student", "bob@example.com", 6.0, [6, 4, 3, 5], **kw)
    return ExportBundle(evaluations=[a, b], kpis_by_version={vid: kpis})


def _col_a(ws) -> list:
    return [(c.row, c.value) for row in ws.iter_rows(min_col=1, max_col=1) for c in row if c.value is not None]


def test_student_sheet_shows_why_and_each_evidence_quote_on_its_own_row() -> None:
    ws = _open(_evidence_bundle())["01 Ann Student"]
    col_a = dict(_col_a(ws))
    clarity_row = next(r for r, v in col_a.items() if v == "Clarity")
    # right under the KPI's score row: Why, then Evidence 1..3, each with its quote in the merged text area
    assert col_a[clarity_row + 1] == "Why" and ws.cell(clarity_row + 1, 2).value == "Reason for Clarity"
    assert col_a[clarity_row + 2] == "Evidence 1\nSection 2"
    assert ws.cell(clarity_row + 2, 2).value == "“We shipped the pilot in six weeks.”"
    assert col_a[clarity_row + 3] == "Evidence 2\nAppendix B"
    assert col_a[clarity_row + 4] == "Evidence 3"
    assert ws.cell(clarity_row + 4, 2).value == "“=HYPERLINK(\"http://evil\",\"x\")”"  # quote marks defuse the formula
    assert ws.row_dimensions[clarity_row + 2].outlineLevel == 1  # collapsible group
    assert ws.row_dimensions[clarity_row].outlineLevel == 0
    # the merged text areas are real merges over B:F
    assert any(m.min_row == clarity_row + 2 and m.min_col == 2 and m.max_col == 6 for m in ws.merged_cells.ranges)


def test_student_sheet_header_card_nav_and_rollup() -> None:
    wb = _open(_evidence_bundle(target=7.0))
    ws = wb["01 Ann Student"]
    assert ws["A1"].value.endswith("Ann Student") and ws.sheet_view.showGridLines is False
    assert ws["A3"].value == "◀ Summary" and ws["A3"].hyperlink.location.strip("'").startswith("Summary")
    assert ws["C3"].value == "◀ Leaderboard"
    assert ws.freeze_panes == "A4"
    col_a = dict(_col_a(ws))
    ov = next(r for r, v in col_a.items() if v == "Overall score")
    assert ws.cell(ov, 2).value == 8.0
    assert _fill(ws.cell(ov, 2)) == P.TARGET_BAND_BY_KEY["exceeds"].tint.lstrip("#")  # 8.0 >= 7.7
    assert ws.cell(ov, 3).value.startswith("Exceeds target  ·  +1.00 vs a target of 7.0")
    assert ws.sheet_properties.tabColor.rgb[-6:].upper() == P.TARGET_BAND_BY_KEY["exceeds"].fill.lstrip("#")
    roll = next(r for r, v in col_a.items() if v == "Category roll-up")
    assert [col_a[roll + 2], col_a[roll + 3]] == ["Communication", "Technical"]
    assert ws.cell(roll + 2, 2).value == pytest.approx((8 * 30 + 6 * 20) / 50)


def test_student_sheet_colours_every_kpi_against_the_scorecards_target() -> None:
    ws = _open(_evidence_bundle(target=4.0))["02 Bob Student"]  # Bob: overall 6.0, KPI scores 6/4/3/5
    col_a = dict(_col_a(ws))
    fills = {col_a[r]: _fill(ws.cell(r, 2)) for r in col_a if col_a[r] in ("Clarity", "Listening", "Depth", "Edge cases")}
    expected = {"Clarity": "exceeds", "Listening": "meets", "Depth": "below", "Edge cases": "exceeds"}  # T=4: exceeds >=4.4, meets >=4.0, below >=3.0
    assert fills == {k: P.TARGET_BAND_BY_KEY[v].tint.lstrip("#") for k, v in expected.items()}


def test_long_reasoning_and_evidence_split_across_rows_within_the_row_limit() -> None:
    vid, kpis = make_version()
    long_quote = "evidence sentence number one is fairly long. " * 400  # ~18k chars, cleaned to 8000
    items = [EvidenceItem(long_quote, "Very long source name " * 6)]
    ev = make_eval(vid, kpis, "Long Text", "l@x.com", 7.0, [7] * 4, evidence=items)
    for res in ev.results.values():
        res.reasoning = "reasoning goes on and on. " * 300
    ws = _open(ExportBundle([ev], {vid: kpis}))["01 Long Text"]
    col_a = dict(_col_a(ws))
    assert any(v == "(cont.)" for v in col_a.values())  # continued on following rows
    assert all((ws.row_dimensions[r].height or 0) <= 409 for r in range(1, ws.max_row + 1))
    texts = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, str)]
    assert sum(len(t) for t in texts if t.startswith("evidence sentence") or t.startswith("“evidence")) > 4000  # not lost


def test_evidence_from_every_stored_shape_is_exported() -> None:
    vid, kpis = make_version()
    shapes = [["plain quote"], [{"quote": "obj quote", "section": "S1", "label": "Slide 4"}], {"p1": ["from dict"]},
              "• bullet one\n• bullet two", None]
    evs = [make_eval(vid, kpis, f"Shape {i}", None, 8.0 - i * 0.1, [8] * 4, evidence=normalize_evidence(raw))
           for i, raw in enumerate(shapes)]
    wb = _open(ExportBundle(evs, {vid: kpis}))
    detail = " ".join(_texts(wb["KPI Detail"]))
    for needle in ("plain quote", "obj quote", "Slide 4 · S1", "from dict", "bullet one", "bullet two"):
        assert needle in detail, needle

