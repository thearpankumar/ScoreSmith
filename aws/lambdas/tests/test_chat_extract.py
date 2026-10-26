"""chat_extract: text of chat attachments (PDF / DOCX) for the backend, through the extract-doc function."""
import io
import json

import pytest

from fakes import MemoryStore
from worker import chat_extract, docs

KEY = "chat/42/0b6c7a1e-1111-4a5b-9c9d-123456789abc"


def event(kind="pdf", **extra):
    return {"mode": "chat_extract", "bucket": "b", "key": f"{KEY}.{kind}", "kind": kind, "out_key": f"{KEY}.json",
            "max_chars": 400000, **extra}


def text_page(doc, text):
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return page


def scanned_page(doc):
    from PIL import Image

    buf = io.BytesIO()
    Image.effect_noise((500, 700), 60).convert("RGB").save(buf, "JPEG")
    page = doc.new_page()
    page.insert_image(page.rect, stream=buf.getvalue())


def table_pdf(path):
    """A page of text with a ruled 2-column table (what find_tables detects), then more text."""
    import pymupdf

    doc = pymupdf.open()
    page = text_page(doc, "Refund procedure. Agents must verify the order before approving anything at all.")
    rows = [["Step", "Owner"], ["Verify order", "Agent"], ["Approve | refund", "Lead"]]
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(72 + c * 200, 120 + r * 24, 272 + c * 200, 144 + r * 24)
            page.draw_rect(rect, color=(0, 0, 0), width=1)
            page.insert_text((rect.x0 + 4, rect.y1 - 7), cell, fontsize=11)
    page.insert_text((72, 230), "After the table.")
    doc.save(path)
    return path


def run(store, ev):
    return chat_extract.run_chat_extract(ev, store)


# ----------------------------------------------------------------------------- PDF
def test_pdf_text_in_reading_order_with_a_markdown_table(tmp_path):
    s = MemoryStore()
    s.put_bytes(f"{KEY}.pdf", table_pdf(tmp_path / "t.pdf").read_bytes())
    out = run(s, event())
    assert out["ok"] and out["kind"] == "pdf" and out["pages"] == 1 and not out["needs_native"]
    text = out["text"]
    assert text.startswith("## Page 1\n\nRefund procedure.")
    assert "| Step | Owner |\n| --- | --- |\n| Verify order | Agent |\n| Approve \\| refund | Lead |" in text
    assert text.index("Refund procedure") < text.index("| Step |") < text.index("After the table.")
    assert out["chars"] == len(text) and out["truncated"] is False and out["scanned_pages"] == []
    assert s.get_json(f"{KEY}.json") == out  # the same answer is stored for the backend
    assert not [k for k in s.objects if k.endswith(".jpg")]  # images are never written


def test_mostly_scanned_pdf_needs_native_reading(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    text_page(doc, "Cover page with a real paragraph of text that is long enough to count as a text page. " * 2)
    scanned_page(doc)
    scanned_page(doc)
    doc.save(tmp_path / "s.pdf")
    s = MemoryStore()
    s.put_bytes(f"{KEY}.pdf", (tmp_path / "s.pdf").read_bytes())
    out = run(s, event())
    assert out["ok"] and out["pages"] == 3 and out["scanned_pages"] == [2, 3] and out["needs_native"] is True
    assert any("2 of 3 page(s)" in w for w in out["warnings"])


def test_one_scanned_page_of_three_is_not_enough_and_blank_pages_are_not_scans(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    text_page(doc, "First page of an SOP with plenty of real text on it, enough to count as a page of text. " * 2)
    scanned_page(doc)
    doc.new_page()  # blank: no image, so not a scan
    doc.save(tmp_path / "m.pdf")
    out = chat_extract.extract_text(str(tmp_path / "m.pdf"), "pdf")
    assert out["scanned_pages"] == [2] and out["needs_native"] is False


def test_long_text_is_truncated_to_max_chars(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    for i in range(5):
        page = doc.new_page()
        for line in range(20):
            page.insert_text((72, 72 + line * 14), f"Page {i} line {line}: lorem ipsum dolor sit amet")
    doc.save(tmp_path / "l.pdf")
    s = MemoryStore()
    s.put_bytes(f"{KEY}.pdf", (tmp_path / "l.pdf").read_bytes())
    out = run(s, event(max_chars=1500))
    assert out["truncated"] is True and len(out["text"]) == out["chars"] == 1500
    assert any("truncated" in w for w in out["warnings"])


def test_the_answer_always_fits_the_synchronous_response_limit():
    big = {"ok": True, "kind": "pdf", "text": "é" * 50_000, "pages": 1, "chars": 50_000, "truncated": False,
           "scanned_pages": [], "needs_native": False, "warnings": []}
    out = chat_extract.fit_response(big, limit=100_000)
    assert len(json.dumps(out)) <= 100_000 and out["truncated"] and out["chars"] == len(out["text"])


# ----------------------------------------------------------------------------- DOCX
def test_docx_text_and_tables_as_markdown(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Escalation policy for the billing desk.")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["KPI", "Target", "Notes"], ["First response", "2 h", "business hours"],
                             ["Resolution", "24 h", ""]]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    t.cell(2, 1).merge(t.cell(2, 2))
    d.add_paragraph("Closing remarks.")
    d.save(tmp_path / "p.docx")
    s = MemoryStore()
    s.put_bytes(f"{KEY}.docx", (tmp_path / "p.docx").read_bytes())
    out = run(s, event("docx"))
    assert out["ok"] and out["kind"] == "docx" and out["pages"] is None and out["needs_native"] is False
    text = out["text"]
    assert "| KPI | Target | Notes |\n| --- | --- | --- |\n| First response | 2 h | business hours |" in text
    assert "| Resolution | 24 h |  |" in text  # the merged cell keeps the grid
    assert text.index("Escalation policy") < text.index("| KPI |") < text.index("Closing remarks.")


def test_evaluation_documents_keep_their_table_rows_format(tmp_path):
    """The evaluation pipeline's content.md is unchanged: a table stays one ' | '-joined line per row."""
    docs.extract_file(str(table_pdf(tmp_path / "t.pdf")), "pdf", tmp_path / "out")
    md = (tmp_path / "out" / "content.md").read_text(encoding="utf-8")
    assert "Step | Owner\nVerify order | Agent" in md and "| --- |" not in md


# ----------------------------------------------------------------------------- bad input never raises
@pytest.mark.parametrize("change, reason", [
    ({"key": "uploads/42/x.pdf"}, "key must be a chat/"),
    ({"key": "chat/42/../../derived/e1/a.pdf"}, "key must be a chat/"),
    ({"key": f"{KEY}.docx"}, "key must be a chat/"),            # extension does not match the kind
    ({"out_key": "derived/e1/corpus.json"}, "out_key must be"),
    ({"out_key": f"{KEY}.txt"}, "out_key must be"),
    ({"kind": "xlsx"}, "kind must be"),
    ({"bucket": ""}, "bucket is missing"),
])
def test_invalid_events_are_refused_without_writing(change, reason):
    s = MemoryStore()
    out = run(s, {**event(), **change})
    assert out["ok"] is False and reason in out["error"] and s.objects == {}


def test_missing_and_corrupt_files_are_answers_not_errors(tmp_path):
    s = MemoryStore()
    out = run(s, event())
    assert out == {"ok": False, "error": "file not found"} and s.get_json(f"{KEY}.json") == out
    s.put_bytes(f"{KEY}.pdf", b"<html>not a pdf</html>")
    out = run(s, event())
    assert out["ok"] is False and "not a valid PDF" in out["error"]
    s.put_bytes(f"{KEY}.docx", b"PK\x03\x04 not really a zip")
    assert run(s, event("docx"))["ok"] is False


def test_the_extract_doc_handler_routes_chat_extract(monkeypatch):
    from worker import handlers

    store = MemoryStore()
    monkeypatch.setattr(chat_extract, "S3Store", lambda bucket: store)
    out = handlers.extract_doc(event())
    assert out == {"ok": False, "error": "file not found"} and store.exists(f"{KEY}.json")
